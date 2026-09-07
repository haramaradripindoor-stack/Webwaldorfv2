from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re

doc = Document()

# Formal Styling
style = doc.styles['Normal']
font = style.font
font.name = 'Helvetica'
font.size = Pt(11)

title = doc.add_heading('Acuerdo de Corresponsabilidad y Convivencia', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('Colegio Waldorf Trekan, Puerto Varas')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True

doc.add_paragraph()

quote = doc.add_paragraph('"La escuela es un lugar de encuentro entre dos generaciones, un espacio de reflexión donde el orden establecido, más que perpetuarse, puede renovarse, acogiendo los nuevos impulsos evolutivos de cada nueva generación."')
quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
quote.runs[0].italic = True
author = doc.add_paragraph('— Rudolf Steiner')
author.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()

sections = [
    ("I. El Impulso Fundacional: Un Refugio Evolutivo, no un Servicio Comercial", 
     "El Colegio Waldorf Trekan no es una empresa prestadora de servicios educativos, ni las familias que lo integran son clientes. Somos un Organismo Vivo. Nuestra comunidad se fundamenta en la Pedagogía Waldorf porque responde a la necesidad biológica y espiritual de proteger la infancia en un mundo hiperacelerado.\n\nEducamos para la libertad, pero —tal como advierte el experto en gestión antroposófica Valentín Wember— comprendemos profundamente que la verdadera libertad nace del orden, el ritmo y la autoridad amorosa del adulto.\n\nPara sostener este refugio, requerimos que cada familia que cruza nuestras puertas asuma un compromiso inquebrantable con la coherencia entre el hogar y la escuela. Lo que más modela el alma de un niño no es el discurso del adulto, sino el esfuerzo visible que hace su comunidad por ir al encuentro del otro, dejando en pausa las conveniencias personales en pos del bienestar superior del niño."),
     
    ("II. La Estructura del Organismo Vivo (Trimembración vs. Asambleísmo)",
     "Valentín Wember es enfático al señalar que una escuela Waldorf sana no es una democracia asambleísta ni una cooperativa de opiniones, sino una república de competencias claras (Trimembración Social). Para evitar el caos institucional y proteger a los maestros, Trekan se organiza en tres esferas autónomas:\n\n1. La Esfera Pedagógica (El Claustro de Maestros): Son la única autoridad en el diseño, ejecución del currículum y manejo del aula, guiados por el estudio profundo del niño. Los apoderados no interfieren en la esfera pedagógica.\n2. La Esfera Administrativa (Directorio/Gestión): Custodian la viabilidad legal, financiera (Economía Fraterna) y de infraestructura del colegio.\n3. La Esfera Comunitaria (Familias y Comisiones): Es el motor vital que sostiene física y anímicamente a la escuela a través de las comisiones de trabajo. Aquí las familias aportan su voluntad y profesionalismo para embellecer y sustentar el colegio."),
     
    ("III. El Sostén del Ritmo: Compromiso de las Familias",
     "Para que el niño tenga sus fuerzas vitales a disposición del aprendizaje, requiere un entorno rítmico y cálido. Al ingresar a Trekan, los apoderados se comprometen irrevocablemente a:\n\nA. Transparencia Absoluta y Cuidado del Aula (Diagnósticos y NEE)\nLa escuela Waldorf es terapéutica por naturaleza, pero no es una clínica de educación especial (educación curativa). Un aula solo puede sostener un porcentaje mínimo de neurodivergencias antes de que el maestro sufra 'burnout' y el ritmo del curso colapse.\n- Las familias se comprometen a transparentar toda la información clínica, conductual y emocional del niño desde el primer día (informes psicológicos, diagnósticos TEA/TDAH o procesos familiares).\n- La omisión intencional de diagnósticos o Necesidades Educativas Especiales (NEE) durante la postulación es motivo de cancelación inmediata de la matrícula.\n- El colegio exige, por contrato, que los niños que requieran apoyo externo (Fonoaudiólogo, Terapeuta Ocupacional) mantengan y cumplan sus terapias como condición de permanencia.\n\nB. El Respeto Sagrado por el Ritmo (Puntualidad y Asistencia)\nLa puntualidad no es una norma burocrática; es el marco de seguridad neurológica del niño. Los atrasos destruyen el momento de mayor conexión rítmica (la ronda matutina) y debilitan tanto al niño que llega tarde como al grupo que es interrumpido.\n- Se exige una asistencia mínima del 80% anual. El incumplimiento injustificado compromete la promoción al grado superior.\n- Protocolo de Atrasos: Si un niño llega con más de 20 minutos de retraso, el apoderado deberá esperar con él fuera del aula hasta el cambio de clase, sin interrumpir al maestro.\n\nC. Protección Sensorial: Pantallas, Vestimenta y Alimentación\n- Cero Pantallas en la Infancia: La exposición a smartphones, tablets y videojuegos altera profundamente la neurobiología infantil, atrofiando su capacidad de asombro y de juego libre. El uso de pantallas está estrictamente restringido en el hogar, y los teléfonos celulares están absolutamente prohibidos dentro del colegio. Cualquier dispositivo será requisado.\n- Vestimenta Cobijadora: Los niños deben asistir con ropa adecuada al clima (botas, trajes de agua, lana) para resguardar su calor vital. Se prohíbe el uso de prendas con estampados de películas, superhéroes o dibujos corporativos, para proteger la pureza del imaginario infantil.\n- Alimentación Vital: Las colaciones deben ser alimentos nobles, sanos y sin procesar, libres de envoltorios plásticos y azúcares refinados."),
     
    ("IV. El Sostén Pedagógico: Compromiso de los Maestros",
     "Nuestros maestros asumen la vocación de ser autoridades amorosas. Se comprometen a:\n- Formarse continuamente a la luz de la Antroposofía y el desarrollo humano.\n- Velar por el bienestar físico, anímico y espiritual de cada niño, ejecutando el currículum Waldorf con rigor.\n- Mantener una comunicación fluida con las familias a través de reuniones individuales y un informe anual profundo sobre el proceso de encarnación del niño.\n- Sugerir derivaciones a especialistas médicos o terapéuticos externos cuando el desarrollo del niño exceda las capacidades del aula común."),
     
    ("V. Arquitectura de la Resolución de Conflictos (Madurez Adulta)",
     "Wember nos recuerda que el mayor daño a una comunidad escolar proviene de los comentarios de pasillo y los grupos de WhatsApp de apoderados. En Trekan, exigimos la máxima madurez en la resolución de diferencias:\n1. Comunicación Directa y Vertical: Cualquier inquietud, duda pedagógica o conflicto debe conversarse en primera instancia, de frente y de manera privada, directamente con el maestro de clase.\n2. Escalamiento Sano: Si la situación no se resuelve en esa instancia, el maestro y la familia elevarán el caso al Equipo Pedagógico (Coordinación).\n3. Resguardo del Vínculo: Si un niño reporta un conflicto en casa, los padres deben informar al maestro al día siguiente para no dar espacio a fantasías o malentendidos. Los padres nunca deben reprender o triangular a niños ajenos.\n\nSi un alumno presenta comportamientos que pongan en riesgo la integridad física o anímica de sí mismo o de otros, el maestro contactará a los padres para el retiro inmediato de la jornada, citando a una reunión de reevaluación del compromiso.")
]

for title_text, body in sections:
    doc.add_heading(title_text, level=1)
    for p in body.split('\n'):
        if p.strip() == '':
            continue
        p_obj = doc.add_paragraph()
        if p.startswith('- '):
            p_obj.style = 'List Bullet'
            p_obj.add_run(p[2:])
        elif re.match(r'^[A-Z0-9]\.', p):
            p_obj.style = 'List Paragraph'
            run = p_obj.add_run(p.split(':', 1)[0] + ':')
            run.bold = True
            if len(p.split(':', 1)) > 1:
                p_obj.add_run(p.split(':', 1)[1])
        else:
            p_obj.add_run(p)
            p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph()
conclusion = doc.add_paragraph()
run = conclusion.add_run('Bienvenidos a Colegio Waldorf Trekan.')
run.bold = True
conclusion.add_run(' Al firmar este documento, sumamos nuestras fuerzas para la conformación de un nuevo orden social, asumiendo la corresponsabilidad de proteger la infancia hoy, para que ellos puedan transformar el mundo mañana.')
conclusion.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('\n\n\n')
signatures = doc.add_table(rows=1, cols=2)
signatures.allow_autofit = True
cell1 = signatures.cell(0, 0)
cell1.text = '_________________________________\nFirma Apoderado 1\nRUT:'
cell2 = signatures.cell(0, 1)
cell2.text = '_________________________________\nFirma Apoderado 2\nRUT:'

import os
downloads_path = os.path.expanduser('~/Downloads/Acuerdo_Corresponsabilidad_Trekan.docx')
doc.save(downloads_path)
print(f"Saved to {downloads_path}")
